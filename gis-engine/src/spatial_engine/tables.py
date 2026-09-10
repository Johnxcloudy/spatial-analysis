from __future__ import annotations

import csv
import json
import math
import re
import sqlite3
import uuid
from contextlib import closing, contextmanager
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, time, timedelta
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

import pyarrow as pa
import pyogrio
import shapely
from openpyxl import load_workbook
from pyproj import CRS, Transformer

from .errors import DomainError, InvalidParamsError
from .validation import require_crs, require_exact_keys, require_path, require_string
from . import vectors

MAX_ROWS = 100_000
MAX_FIELDS = 256
MAX_CELLS = 2_000_000
MAX_VALUE_BYTES = 128 * 1024 * 1024
MAX_SOURCE_BYTES = 128 * 1024 * 1024
MAX_ZIP_EXPANDED_BYTES = 256 * 1024 * 1024
MAX_ZIP_MEMBERS = 10_000
MAX_CELL_CHARS = 65_536
MAX_PREVIEW_ROWS = 20
MAX_PREVIEW_BYTES = 1024 * 1024
MAX_SHEETS = 512
MAX_HEADER_ROW = 1000
OPTION_KEYS = {"sourcePath", "encoding", "delimiter", "sheet", "headerRow"}
CELL_TYPES = {"n", "s", "b", "d", "f", "e", "inlineStr", "str"}
VALUE_TYPES = {"null", "string", "int", "float", "bool", "date", "datetime", "time", "duration", "formula", "error"}
COORDINATE_NUMBER = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?")


def validate_options(params: dict, *, for_import: bool = False) -> dict:
    require_exact_keys(params, OPTION_KEYS)
    path = require_path(params["sourcePath"], "sourcePath")
    if path.suffix.lower() not in {".csv", ".xlsx"} or not path.is_file():
        raise DomainError("Source must be an existing CSV or XLSX file", kind="unsupported_source")
    if path.stat().st_size > MAX_SOURCE_BYTES:
        raise DomainError("Table source exceeds the file size limit", kind="source_limit")
    header = params["headerRow"]
    if isinstance(header, bool) or not isinstance(header, int) or not 1 <= header <= MAX_HEADER_ROW:
        raise InvalidParamsError("headerRow must be a one-based row from 1 to 1000")
    sheet = params["sheet"]
    if sheet is not None:
        sheet = require_string(sheet, "sheet", maximum=256)
    if path.suffix.lower() == ".csv":
        if sheet is not None:
            raise InvalidParamsError("CSV sheet must be null")
        encoding = require_string(params["encoding"], "encoding", maximum=32)
        if encoding.upper().replace("-", "") not in {"UTF8", "UTF8SIG", "GBK", "GB18030"}:
            raise InvalidParamsError("encoding must be UTF-8, UTF-8-sig, GBK or GB18030")
        delimiter = params["delimiter"]
        if not isinstance(delimiter, str) or delimiter not in {",", ";", "\t", "|"}:
            raise InvalidParamsError("delimiter must be comma, semicolon, tab or pipe")
    else:
        encoding, delimiter = None, None
        if for_import and sheet is None:
            raise InvalidParamsError("Select an XLSX sheet explicitly before importing")
    return {"sourcePath": str(path), "encoding": encoding, "delimiter": delimiter, "sheet": sheet, "headerRow": header}


def _text(value: str) -> str:
    if len(value) > MAX_CELL_CHARS:
        raise DomainError("Table cell exceeds the text size limit", kind="source_limit")
    if "\x00" in value:
        raise DomainError("Embedded NUL text cannot be preserved in a GeoPackage field", kind="unsupported_fields")
    try:
        value.encode("utf-8", errors="strict")
    except UnicodeError as exc:
        raise DomainError("Table text cannot be decoded without replacement", kind="encoding_error") from exc
    return value


def _normalize_cell(cell) -> tuple[str | None, dict | None]:
    value, cell_type = cell.value, cell.data_type
    if cell_type not in CELL_TYPES:
        raise DomainError("Unsupported XLSX cell type", kind="unsupported_fields", detail=str(cell_type))
    if value is None:
        normalized, value_type = None, "null"
    elif cell_type == "f":
        if not isinstance(value, str):
            raise DomainError("Unsupported XLSX formula representation", kind="unsupported_fields")
        normalized, value_type = value, "formula"
    elif cell_type == "e":
        normalized, value_type = str(value), "error"
    elif isinstance(value, bool):
        normalized, value_type = "true" if value else "false", "bool"
    elif isinstance(value, datetime):
        normalized, value_type = value.isoformat(), "datetime"
    elif isinstance(value, date):
        normalized, value_type = value.isoformat(), "date"
    elif isinstance(value, time):
        normalized, value_type = value.isoformat(), "time"
    elif isinstance(value, timedelta):
        normalized, value_type = str(value), "duration"
    elif isinstance(value, int):
        normalized, value_type = str(value), "int"
    elif isinstance(value, float):
        if not math.isfinite(value):
            raise DomainError("Non-finite XLSX numeric cells are unsupported", kind="unsupported_fields")
        normalized, value_type = repr(value), "float"
    elif isinstance(value, str):
        normalized, value_type = value, "string"
    else:
        raise DomainError("Unsupported XLSX scalar value", kind="unsupported_fields", detail=type(value).__name__)
    if normalized is not None:
        normalized = _text(normalized)
    number_format = _text(cell.number_format or "General")
    metadata = None
    if value_type not in {"string", "null"} or number_format != "General":
        metadata = {"cell_type": cell_type, "value_type": value_type, "number_format": number_format}
    return normalized, metadata


def _check_xlsx(path: Path) -> None:
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_MEMBERS or sum(item.file_size for item in members) > MAX_ZIP_EXPANDED_BYTES:
                raise DomainError("XLSX expanded content exceeds the resource limit", kind="source_limit")
            if any(item.flag_bits & 1 for item in members):
                raise DomainError("Encrypted XLSX content is unsupported", kind="unsupported_source")
            if len({item.filename for item in members}) != len(members):
                raise DomainError("XLSX contains duplicate archive members", kind="invalid_source")
    except BadZipFile as exc:
        raise DomainError("XLSX is not a readable workbook archive", kind="source_read_failed") from exc


@contextmanager
def _source_rows(options: dict, cancelled):
    path = Path(options["sourcePath"])
    if path.suffix.lower() == ".csv":
        encoding = options["encoding"]
        codec = "utf-8-sig" if encoding.upper().replace("-", "") in {"UTF8", "UTF8SIG"} else encoding
        prior_limit = csv.field_size_limit(MAX_CELL_CHARS)
        try:
            with path.open("r", encoding=codec, errors="strict", newline="") as handle:
                reader = csv.reader(handle, delimiter=options["delimiter"], strict=True)

                def rows():
                    for number, row in enumerate(reader, 1):
                        vectors._cancel(cancelled)
                        yield number, [_text(value) for value in row], {}

                yield rows(), [], None, {"sourceRowPolicy": "one-based CSV logical record"}
        finally:
            csv.field_size_limit(prior_limit)
        return
    _check_xlsx(path)
    vectors._cancel(cancelled)
    source_handle = path.open("rb")
    workbook = None
    try:
        workbook = load_workbook(source_handle, read_only=True, data_only=False, keep_links=False)
        sheets = workbook.sheetnames
        if not sheets or len(sheets) > MAX_SHEETS:
            raise DomainError("XLSX sheet count is outside the supported limit", kind="source_limit")
        selected = options["sheet"] or sheets[0]
        if selected not in sheets:
            raise DomainError("Selected XLSX sheet does not exist", kind="sheet_not_found")
        sheet = workbook[selected]
        sheet.reset_dimensions()

        def rows():
            for number, row in enumerate(sheet.iter_rows(), 1):
                vectors._cancel(cancelled)
                if len(row) + 2 > MAX_FIELDS:
                    raise DomainError("Table exceeds the field limit including internal IDs", kind="source_limit")
                values, metadata = [], {}
                for index, cell in enumerate(row):
                    value, details = _normalize_cell(cell)
                    values.append(value)
                    if details is not None:
                        metadata[index] = details
                yield number, values, metadata

        yield rows(), sheets, selected, {
            "sourceRowPolicy": "one-based worksheet row", "workbookEpoch": workbook.epoch.isoformat(),
            "formulaPolicy": "retain_text_no_execution",
            "normalizationPolicy": "scalar values to text/NULL; sparse scalar types and number formats retained",
        }
    finally:
        if workbook is not None:
            workbook.close()
        source_handle.close()


@dataclass
class _Scan:
    columns: list[dict] = field(default_factory=list)
    rows: list[tuple[int, list[str | None]]] = field(default_factory=list)
    cells: list[dict] = field(default_factory=list)
    sheets: list[str] = field(default_factory=list)
    sheet: str | None = None
    metadata: dict[str, str] = field(default_factory=dict)
    source_cells: int = 0
    value_bytes: int = 0
    truncated: bool = False


def _columns(headers: list[str | None]) -> list[dict]:
    if not headers:
        raise DomainError("Selected header row has no columns", kind="invalid_header")
    if len(headers) + 2 > MAX_FIELDS:
        raise DomainError("Table exceeds the field limit including internal IDs", kind="source_limit")
    names, result = set(), []
    for index, header in enumerate(headers, 1):
        base = (header or "").strip()
        base = "".join(character if ord(character) >= 32 else "_" for character in base)[:128]
        name = vectors._unique_name(base or f"column_{index}", names)
        result.append({"index": index, "sourceName": header, "fieldName": name})
    return result


def _scan(options: dict, *, preview: bool, progress, cancelled) -> _Scan:
    result, headers = _Scan(), None
    try:
        with _source_rows(options, cancelled) as (rows, sheets, selected, metadata):
            result.sheets, result.sheet, result.metadata = list(sheets), selected, metadata
            for source_row, values, cell_metadata in rows:
                result.source_cells += len(values)
                result.value_bytes += sum(len(value.encode("utf-8")) for value in values if value is not None)
                result.value_bytes += sum(len(json.dumps(value, ensure_ascii=False).encode("utf-8")) for value in cell_metadata.values())
                if result.source_cells > MAX_CELLS or result.value_bytes > MAX_VALUE_BYTES:
                    raise DomainError("Table exceeds the source cell or normalized byte limit", kind="source_limit")
                if source_row < options["headerRow"]:
                    continue
                if source_row == options["headerRow"]:
                    headers = values
                    result.columns = _columns(headers)
                    continue
                if headers is None:
                    raise DomainError("Selected header row was not found", kind="invalid_header")
                if preview and len(result.rows) >= MAX_PREVIEW_ROWS:
                    result.truncated = True
                    break
                if len(result.rows) >= MAX_ROWS:
                    raise DomainError("Table exceeds the data row limit", kind="source_limit")
                if options["sourcePath"].lower().endswith(".csv") and len(values) != len(headers):
                    raise DomainError("CSV record width does not match its header", kind="invalid_table_row", detail=f"record {source_row}")
                if (len(result.rows) + 1) * max(len(headers), len(values)) > MAX_CELLS:
                    raise DomainError("Normalized table rectangle exceeds the cell limit", kind="source_limit")
                if len(values) > len(headers):
                    extra = len(values) - len(headers)
                    headers.extend([None] * extra)
                    result.columns = _columns(headers)
                    for _, preceding in result.rows:
                        preceding.extend([None] * extra)
                values.extend([None] * (len(headers) - len(values)))
                row_id = str(len(result.rows) + 1)
                result.rows.append((source_row, values))
                for index, details in cell_metadata.items():
                    result.cells.append({"row_id": row_id, "field_name": result.columns[index]["fieldName"], **details})
                if len(result.rows) % vectors.BATCH_SIZE == 0:
                    progress("reading", len(result.rows), None)
            if headers is None:
                raise DomainError("Selected header row was not found", kind="invalid_header")
    except (DomainError, InvalidParamsError):
        raise
    except UnicodeError as exc:
        raise DomainError("Table text could not be decoded with the selected encoding", kind="encoding_error", detail=str(exc)) from exc
    except Exception as exc:
        raise DomainError("Could not read table source", kind="source_read_failed", detail=str(exc)) from exc
    return result


def inspect_table(params: dict) -> dict:
    options = validate_options(params)
    if options["sourcePath"].lower().endswith(".xlsx") and options["sheet"] is None:
        path = Path(options["sourcePath"])
        _check_xlsx(path)
        try:
            with path.open("rb") as source_handle:
                workbook = load_workbook(source_handle, read_only=True, data_only=False, keep_links=False)
                try:
                    sheets = workbook.sheetnames
                    if not sheets or len(sheets) > MAX_SHEETS:
                        raise DomainError("XLSX sheet count is outside the supported limit", kind="source_limit")
                    return {"sourcePath": str(path), "driver": "XLSX", "sheets": sheets, "sheet": None,
                            "columns": [], "rows": [], "truncated": False, "warnings": []}
                finally:
                    workbook.close()
        except DomainError:
            raise
        except Exception as exc:
            raise DomainError("Could not inspect XLSX workbook", kind="source_read_failed", detail=str(exc)) from exc
    scan = _scan(options, preview=True, progress=lambda *args: None, cancelled=lambda: False)
    columns = [{**column, "sourceName": column["sourceName"][:256] if column["sourceName"] is not None else None}
               for column in scan.columns]
    warnings = ["Preview is partial; complete validation runs during import."]
    if any(column["sourceName"] is not None and len(column["sourceName"]) > 256 for column in scan.columns):
        warnings.append("Long headers are shortened in preview; their complete values are stored in column_metadata.")
    if scan.sheet is not None:
        warnings.append("XLSX values are normalized text/NULL; formulas are retained as text and never executed.")
    result = {"sourcePath": options["sourcePath"], "driver": "XLSX" if scan.sheet is not None else "CSV",
              "sheets": scan.sheets, "sheet": scan.sheet, "columns": columns,
              "rows": [{"sourceRow": number, "values": {column["fieldName"]: value for column, value in zip(scan.columns, values)}}
                       for number, values in scan.rows], "truncated": scan.truncated, "warnings": warnings}
    while len(json.dumps(result, ensure_ascii=True, separators=(",", ":")).encode("ascii")) > MAX_PREVIEW_BYTES:
        if not result["rows"]:
            raise DomainError("Table preview columns exceed the response limit", kind="source_limit")
        result["rows"].pop()
        result["truncated"] = True
    return result


def _dataset_id(value: Any) -> str:
    result = require_string(value, "datasetId", maximum=64)
    try:
        uuid.UUID(result)
    except ValueError as exc:
        raise InvalidParamsError("datasetId must be a UUID") from exc
    return result


def _artifact(work_dir: Path) -> Path:
    work_dir.mkdir(parents=True, exist_ok=True)
    path = (work_dir / "snapshot.gpkg").resolve()
    if path.exists():
        raise DomainError("Snapshot destination already exists", kind="destination_exists")
    return path


def _auxiliary(rows: list[dict], fields: dict[str, pa.DataType]) -> pa.Table:
    return pa.table({name: pa.array([row[name] for row in rows], type=kind) for name, kind in fields.items()})


def import_table(payload: dict, work_dir: Path, *, progress, cancelled) -> dict:
    require_exact_keys(payload, OPTION_KEYS | {"datasetId"})
    options = validate_options({key: payload[key] for key in OPTION_KEYS}, for_import=True)
    dataset_id = _dataset_id(payload["datasetId"])
    vectors._cancel(cancelled)
    source = Path(options["sourcePath"])
    progress("fingerprinting", None, None)
    fingerprint = vectors.source_fingerprint(source, cancelled)
    scan = _scan(options, preview=False, progress=progress, cancelled=cancelled)
    names = {column["fieldName"].casefold() for column in scan.columns}
    id_field = vectors._unique_name("_sa_id", names)
    source_field = vectors._unique_name("_sa_source_row", names)
    table = pa.table({column["fieldName"]: pa.array([row[index] for _, row in scan.rows], type=pa.string())
                      for index, column in enumerate(scan.columns)})
    table = table.append_column(id_field, pa.array([str(index + 1) for index in range(len(scan.rows))], type=pa.string()))
    table = table.append_column(source_field, pa.array([str(number) for number, _ in scan.rows], type=pa.string()))
    column_table = _auxiliary(
        [{"source_index": column["index"], "source_name": column["sourceName"], "field_name": column["fieldName"]} for column in scan.columns],
        {"source_index": pa.int64(), "source_name": pa.string(), "field_name": pa.string()},
    )
    auxiliary = {"column_metadata": column_table}
    if scan.sheet is not None:
        auxiliary["cell_metadata"] = _auxiliary(scan.cells, {
            "row_id": pa.string(), "field_name": pa.string(), "cell_type": pa.string(),
            "value_type": pa.string(), "number_format": pa.string(),
        })
    fields = []
    for column in scan.columns:
        metadata = vectors._field(column["fieldName"], "CSV text" if scan.sheet is None else "XLSX normalized scalar", "string", True)
        if column["sourceName"] and len(column["sourceName"]) <= 1024:
            metadata["alias"] = column["sourceName"]
        fields.append(metadata)
    warnings = ["Values are preserved according to the recorded parser options; no column type inference was applied."]
    if scan.sheet is not None:
        warnings.append("XLSX displayed number formatting is not applied to stored text; scalar types and format codes are retained separately.")
    metadata = {**scan.metadata, "headerRow": str(options["headerRow"]), "columnMetadataLayer": "column_metadata",
                "encoding": options["encoding"] or "not applicable", "delimiter": options["delimiter"] or "not applicable"}
    dataset = {"id": dataset_id, "version": "0" * 64, "name": (scan.sheet or source.stem)[:200], "kind": "table",
               "source": {"path": str(source), "layer": scan.sheet or source.stem, "driver": "XLSX" if scan.sheet else "CSV",
                          "fingerprint": fingerprint, "encoding": options["encoding"], "assignedCrs": None, "crsWkt": None, "metadata": metadata},
               "relativePath": f"datasets/{dataset_id}.gpkg", "storageLayer": "records", "featureCount": len(scan.rows),
               "geometryType": None, "crsWkt": None, "crsAuthority": None, "bounds": None, "boundsWgs84": None,
               "fields": fields, "internalIdField": id_field, "sourceFidField": source_field,
               "cellMetadataLayer": "cell_metadata" if scan.sheet is not None else None,
               "report": {"status": "warning", "checks": [], "warnings": warnings,
                          "notChecked": ["Workbook layout, rich-text formatting, formula recalculation and original XML numeric representation", "Coordinate semantics and source positional accuracy"],
                          "counts": {"rows": len(scan.rows), "cells": scan.source_cells, "normalizedBytes": scan.value_bytes,
                                     "cellMetadataRows": len(scan.cells), "formulaCells": sum(row["cell_type"] == "f" for row in scan.cells)},
                          "validatorVersion": "1"}, "createdAt": datetime.now(UTC).isoformat()}
    vectors._cancel(cancelled)
    artifact = _artifact(work_dir)
    try:
        progress("writing", len(scan.rows), len(scan.rows))
        pyogrio.write_arrow(table, artifact, layer="records", driver="GPKG", layer_options={"FID": vectors._unique_name("_sa_fid", names)})
        for layer, data in auxiliary.items():
            vectors._cancel(cancelled)
            pyogrio.write_arrow(data, artifact, layer=layer, driver="GPKG")
        progress("validating", len(scan.rows), len(scan.rows))
        vectors.verify_snapshot(artifact, dataset, table, cancelled=cancelled)
        verify_table_auxiliary(artifact, dataset, expected=auxiliary, cancelled=cancelled)
        if vectors.source_fingerprint(source, cancelled) != fingerprint:
            raise DomainError("Source changed during table import", kind="source_changed")
        dataset["version"] = vectors._content_hash(artifact, cancelled)
        dataset["report"]["checks"] = [
            {"code": "roundtrip", "passed": True, "detail": "All normalized values, NULLs, IDs, header mappings and scalar metadata reread and compared"},
            {"code": "source_unchanged", "passed": True, "detail": "Source fingerprints match before and after import"},
            {"code": "nonspatial_gpkg", "passed": True, "detail": "GeoPackage attributes tables registered without geometry"},
        ]
        return {"dataset": dataset, "artifactPath": str(artifact)}
    except Exception:
        if artifact.exists():
            artifact.unlink()
        raise


def verify_table_auxiliary(path: Path, dataset: dict, *, expected: dict[str, pa.Table] | None = None, cancelled=lambda: False) -> None:
    required = {"records", "column_metadata"}
    if dataset["cellMetadataLayer"] is not None:
        if dataset["cellMetadataLayer"] != "cell_metadata":
            raise DomainError("Unsupported table metadata layer", kind="invalid_snapshot")
        required.add("cell_metadata")
    with closing(sqlite3.connect(f"{path.resolve().as_uri()}?mode=ro", uri=True)) as connection:
        contents = dict(connection.execute("SELECT table_name, data_type FROM gpkg_contents"))
        if set(contents) != required or any(value != "attributes" for value in contents.values()):
            raise DomainError("Table snapshot does not contain the expected GeoPackage attributes tables", kind="roundtrip_failed")
        if connection.execute("SELECT COUNT(*) FROM gpkg_geometry_columns").fetchone()[0] != 0:
            raise DomainError("Table snapshot unexpectedly contains geometry", kind="roundtrip_failed")
    columns = pyogrio.read_arrow(path, layer="column_metadata")[1]
    if columns.column_names != ["source_index", "source_name", "field_name"]:
        raise DomainError("Table column provenance schema is invalid", kind="invalid_snapshot")
    if columns["field_name"].to_pylist() != [field["name"] for field in dataset["fields"]]:
        raise DomainError("Table column provenance does not match its fields", kind="invalid_snapshot")
    if columns["source_index"].to_pylist() != list(range(1, len(dataset["fields"]) + 1)):
        raise DomainError("Table column positions are invalid", kind="invalid_snapshot")
    if "cell_metadata" in required:
        cells = pyogrio.read_arrow(path, layer="cell_metadata")[1]
        if cells.column_names != ["row_id", "field_name", "cell_type", "value_type", "number_format"]:
            raise DomainError("Table cell provenance schema is invalid", kind="invalid_snapshot")
        if len(cells) > MAX_CELLS or len(cells) != dataset["report"]["counts"]["cellMetadataRows"]:
            raise DomainError("Table cell provenance count is invalid", kind="invalid_snapshot")
        names, seen = {field["name"] for field in dataset["fields"]}, set()
        for row in cells.to_pylist():
            vectors._cancel(cancelled)
            key = (row["row_id"], row["field_name"])
            identifier = row["row_id"]
            if (key in seen or row["field_name"] not in names or not isinstance(identifier, str)
                    or not identifier.isascii() or not identifier.isdecimal() or not 1 <= int(identifier) <= dataset["featureCount"]
                    or row["cell_type"] not in CELL_TYPES or row["value_type"] not in VALUE_TYPES or not isinstance(row["number_format"], str)):
                raise DomainError("Table cell provenance is invalid", kind="invalid_snapshot")
            seen.add(key)
    if expected is not None:
        for layer, table in expected.items():
            vectors._cancel(cancelled)
            actual = pyogrio.read_arrow(path, layer=layer)[1]
            if table.column_names != actual.column_names or table.to_pydict() != actual.to_pydict():
                raise DomainError("Table auxiliary metadata did not round trip", kind="roundtrip_failed", detail=layer)


def table_to_points(payload: dict, work_dir: Path, *, progress, cancelled) -> dict:
    require_exact_keys(payload, {"dataset", "managedPath", "datasetId", "xField", "yField", "declaredCrs"})
    parent = payload["dataset"]
    if not isinstance(parent, dict) or parent.get("kind") != "table":
        raise DomainError("Point generation requires a table dataset", kind="dataset_not_table")
    dataset_id = _dataset_id(payload["datasetId"])
    if dataset_id == parent["id"]:
        raise InvalidParamsError("Point dataset must have a new identity")
    known = {field["name"] for field in parent["fields"]}
    for key in ("xField", "yField"):
        if not isinstance(payload[key], str) or payload[key] not in known:
            raise InvalidParamsError(f"{key} must name a source table field")
    if payload["xField"] == payload["yField"]:
        raise InvalidParamsError("X and Y must use distinct table fields")
    if len(parent["fields"]) + 4 > MAX_FIELDS:
        raise DomainError("Point result exceeds the field limit including IDs and coordinate status fields", kind="source_limit")
    declared = require_crs(payload["declaredCrs"], "declaredCrs")
    crs = CRS.from_user_input(declared)
    if len(crs.axis_info) != 2 or not (crs.is_geographic or crs.is_projected) or crs.is_compound:
        raise DomainError("Coordinates require a two-dimensional geographic or projected CRS", kind="unsupported_crs")
    if crs.is_geographic and any(not math.isclose(axis.unit_conversion_factor, math.pi / 180) for axis in crs.axis_info):
        raise DomainError("Geographic coordinate input currently requires degree units", kind="unsupported_crs")
    managed = require_path(payload["managedPath"], "managedPath")
    vectors._cancel(cancelled)
    progress("validating", None, None)
    vectors.verify_snapshot(managed, parent, cancelled=cancelled)
    table = pyogrio.read_arrow(managed, layer=parent["storageLayer"])[1]
    cell_types = {}
    if parent["cellMetadataLayer"] is not None:
        for row in pyogrio.read_arrow(managed, layer=parent["cellMetadataLayer"])[1].to_pylist():
            if row["field_name"] in {payload["xField"], payload["yField"]}:
                cell_types[(row["row_id"], row["field_name"])] = row["value_type"]
    transformer = Transformer.from_crs(crs, 4326, always_xy=True)
    identifiers = table[parent["internalIdField"]].to_pylist()
    x_values, y_values = table[payload["xField"]].to_pylist(), table[payload["yField"]].to_pylist()
    geometries, statuses, reasons, geographic = [], [], [], []
    for index, (identifier, raw_x, raw_y) in enumerate(zip(identifiers, x_values, y_values)):
        vectors._cancel(cancelled)
        reason, point = None, None
        forbidden = [cell_types.get((identifier, payload[key])) for key in ("xField", "yField")]
        denied = next((kind for kind in forbidden if kind in {"formula", "error", "bool", "date", "datetime", "time", "duration"}), None)
        if denied:
            reason = f"unsupported_{denied}_coordinate"
        elif raw_x is None or raw_y is None or not raw_x.strip() or not raw_y.strip():
            reason = "missing_coordinate"
        elif COORDINATE_NUMBER.fullmatch(raw_x.strip()) is None or COORDINATE_NUMBER.fullmatch(raw_y.strip()) is None:
            reason = "invalid_numeric_coordinate"
        else:
            try:
                x, y = float(raw_x.strip()), float(raw_y.strip())
                if not math.isfinite(x) or not math.isfinite(y):
                    reason = "non_finite_coordinate"
                elif crs.is_geographic and not (-180 <= x <= 180 and -90 <= y <= 90):
                    reason = "geographic_coordinate_out_of_range"
                else:
                    longitude, latitude = transformer.transform(x, y, errcheck=True)
                    if not math.isfinite(longitude) or not math.isfinite(latitude) or not (-180 <= longitude <= 180 and -90 <= latitude <= 90):
                        reason = "coordinate_outside_transform_domain"
                    else:
                        point = shapely.Point(x, y)
                        geographic.append((longitude, latitude))
            except (ValueError, OverflowError):
                reason = "invalid_numeric_coordinate"
            except Exception:
                reason = "coordinate_outside_transform_domain"
        geometries.append(point)
        statuses.append("valid" if point is not None else "invalid")
        reasons.append(reason)
        if (index + 1) % vectors.BATCH_SIZE == 0:
            progress("deriving", index + 1, len(table))
    valid = len(geographic)
    if valid == 0:
        raise DomainError("No rows contain valid coordinates in the declared CRS", kind="no_valid_coordinates")
    names = {name.casefold() for name in table.column_names}
    status_field = vectors._unique_name("_sa_coordinate_status", names)
    reason_field = vectors._unique_name("_sa_coordinate_reason", names)
    geometry_field = vectors._unique_name("_sa_geometry", names)
    table = table.append_column(status_field, pa.array(statuses, type=pa.string()))
    table = table.append_column(reason_field, pa.array(reasons, type=pa.string()))
    table = table.append_column(geometry_field, pa.array(shapely.to_wkb(geometries).tolist(), type=pa.binary()))
    bounds_array = shapely.bounds([geometry for geometry in geometries if geometry is not None])
    bounds = [float(bounds_array[:, 0].min()), float(bounds_array[:, 1].min()), float(bounds_array[:, 2].max()), float(bounds_array[:, 3].max())]
    crs_wkt, authority = vectors._crs(crs)
    fields = [dict(value) for value in parent["fields"]] + [vectors._field(status_field, "coordinate status", "string", True),
                                                           vectors._field(reason_field, "coordinate validation reason", "string", True)]
    metadata = {"parentDatasetId": parent["id"], "parentVersion": parent["version"], "xField": payload["xField"], "yField": payload["yField"],
                "declaredCrs": declared, "coordinateOrder": "X=longitude/easting;Y=latitude/northing", "invalidRowPolicy": "retain_NULL_geometry",
                "numericParsingPolicy": "ASCII decimal/scientific notation; outer whitespace trimmed; no locale or thousands separators",
                "coordinateStatusField": status_field, "coordinateReasonField": reason_field}
    dataset = {"id": dataset_id, "version": "0" * 64, "name": f"{parent['name'][:193]} points", "kind": "vector",
               "source": {"path": str(managed), "layer": parent["storageLayer"], "driver": "TablePoints", "fingerprint": parent["version"],
                          "encoding": None, "assignedCrs": declared, "crsWkt": None, "metadata": metadata},
               "relativePath": f"datasets/{dataset_id}.gpkg", "storageLayer": "features", "featureCount": len(table), "geometryType": "Point",
               "crsWkt": crs_wkt, "crsAuthority": authority, "bounds": bounds,
               "boundsWgs84": [min(item[0] for item in geographic), min(item[1] for item in geographic), max(item[0] for item in geographic), max(item[1] for item in geographic)],
               "fields": fields, "internalIdField": parent["internalIdField"], "sourceFidField": parent["sourceFidField"],
               "report": {"status": "restricted" if valid != len(table) else "warning", "checks": [],
                          "warnings": ["CRS was explicitly declared; successful conversion does not verify source position or suitability for area analysis."],
                          "notChecked": ["Source positional accuracy", "Coordinate semantics beyond the selected fields and declared CRS"],
                          "counts": {"features": len(table), "validCoordinates": valid, "invalidCoordinates": len(table) - valid,
                                     "missing": len(table) - valid, "vertices": valid, "invalid": 0, "empty": 0}, "validatorVersion": "1"},
               "createdAt": datetime.now(UTC).isoformat()}
    artifact = _artifact(work_dir)
    try:
        vectors._cancel(cancelled)
        progress("writing", len(table), len(table))
        pyogrio.write_arrow(table, artifact, layer="features", driver="GPKG", geometry_name=geometry_field, geometry_type="Point", crs=crs_wkt,
                            layer_options={"SPATIAL_INDEX": "YES", "GEOMETRY_NAME": geometry_field, "FID": vectors._unique_name("_sa_fid", names)})
        progress("validating", len(table), len(table))
        vectors.verify_snapshot(artifact, dataset, table, geometry_field, cancelled)
        if vectors._content_hash(managed, cancelled) != parent["version"]:
            raise DomainError("Source table changed during point generation", kind="source_changed")
        dataset["version"] = vectors._content_hash(artifact, cancelled)
        dataset["report"]["checks"] = [
            {"code": "roundtrip", "passed": True, "detail": "Attributes, row IDs, point coordinates and NULL geometries reread and compared"},
            {"code": "parent_unchanged", "passed": True, "detail": "Parent table content hash matches its immutable version"},
            {"code": "coordinates_valid", "passed": valid == len(table), "detail": "Invalid rows retained with reasons and NULL geometry", "count": len(table) - valid},
        ]
        return {"dataset": dataset, "artifactPath": str(artifact)}
    except Exception:
        if artifact.exists():
            artifact.unlink()
        raise
