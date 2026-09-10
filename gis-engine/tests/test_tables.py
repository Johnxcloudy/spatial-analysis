from __future__ import annotations

import csv
import gc
import importlib
import sqlite3
import uuid
from contextlib import closing
from pathlib import Path

import pyogrio
import pytest
import shapely
from openpyxl import Workbook

from spatial_engine.errors import DomainError, InvalidParamsError
from spatial_engine import vector_queries, vectors
from table_fixtures import create_table_fixture_bundle


def tables():
    return importlib.import_module("spatial_engine.tables")


def options(path, **changes):
    csv_source = Path(path).suffix.lower() == ".csv"
    result = {"sourcePath": str(path), "encoding": "UTF-8" if csv_source else None,
              "delimiter": "," if csv_source else None, "sheet": None, "headerRow": 1}
    result.update(changes)
    return result


def import_table(path, work_dir, **changes):
    return tables().import_table(
        {**options(path, **changes), "datasetId": str(uuid.uuid4())}, work_dir,
        progress=lambda *args: None, cancelled=lambda: False,
    )


def points(result, work_dir, **changes):
    payload = {"dataset": result["dataset"], "managedPath": result["artifactPath"],
               "datasetId": str(uuid.uuid4()), "xField": "x", "yField": "y", "declaredCrs": "EPSG:4326"}
    payload.update(changes)
    return tables().table_to_points(payload, work_dir, progress=lambda *args: None, cancelled=lambda: False)


@pytest.fixture
def sources(tmp_path):
    return create_table_fixture_bundle(tmp_path / "sources")


def test_csv_table_snapshot_preserves_exact_strings_and_source_identity(sources, tmp_path):
    source = Path(sources["csv"])
    original = source.read_bytes()
    result = import_table(source, tmp_path / "work")
    dataset = result["dataset"]
    assert dataset["kind"] == "table" and dataset["featureCount"] == 4
    assert dataset["geometryType"] is None and dataset["crsWkt"] is None
    assert dataset["bounds"] is None and dataset["boundsWgs84"] is None
    assert dataset["cellMetadataLayer"] is None
    stored = pyogrio.read_arrow(result["artifactPath"], layer="records")[1]
    assert stored["code"].to_pylist() == ["001", "002", "003", "004"]
    assert stored["label"].to_pylist() == ["Chinese \u571f\u5730", "NULL", "outside", ""]
    assert stored["x"].to_pylist() == ["114", "", "200", " 114.5 "]
    assert stored[dataset["sourceFidField"]].to_pylist() == ["2", "3", "4", "5"]
    with closing(sqlite3.connect(result["artifactPath"])) as connection:
        assert connection.execute("SELECT data_type FROM gpkg_contents WHERE table_name='records'").fetchone() == ("attributes",)
        assert connection.execute("SELECT COUNT(*) FROM gpkg_geometry_columns").fetchone() == (0,)
    vectors.verify_snapshot(Path(result["artifactPath"]), dataset)
    exported = vectors.export_vector({"dataset": dataset, "managedPath": result["artifactPath"]},
                                     tmp_path / "export", lambda *args: None, lambda: False)
    assert Path(exported["artifactPath"]).read_bytes() == Path(result["artifactPath"]).read_bytes()
    assert source.read_bytes() == original


def test_duplicate_blank_and_long_headers_have_complete_mapping(tmp_path):
    path = tmp_path / "headers.csv"
    headers = ["code", "CODE", "", "_sa_id", "h" * 1500]
    with path.open("w", encoding="utf-8", newline="") as handle:
        csv.writer(handle).writerows([headers, ["001", "002", "blank", "original", "long"]])
    inspection = tables().inspect_table(options(path))
    names = [column["fieldName"] for column in inspection["columns"]]
    assert len(set(name.casefold() for name in names)) == len(headers)
    assert all(len(name) <= 256 for name in names)
    result = import_table(path, tmp_path / "work")
    mappings = pyogrio.read_arrow(result["artifactPath"], layer="column_metadata")[1].to_pydict()
    assert mappings["source_name"] == headers
    assert mappings["field_name"] == names
    assert all(field["alias"] is None or len(field["alias"]) <= 1024 for field in result["dataset"]["fields"])


@pytest.mark.parametrize("encoding,delimiter", [("GBK", ";"), ("GB18030", "\t"), ("UTF-8", ",")])
def test_explicit_csv_encodings_delimiters_and_header_row(tmp_path, encoding, delimiter):
    path = tmp_path / "encoded.csv"
    with path.open("w", encoding=encoding, newline="") as handle:
        writer = csv.writer(handle, delimiter=delimiter)
        writer.writerows([["preamble"], ["code", "name"], ["001", "\u571f\u5730"]])
    result = import_table(path, tmp_path / "work", encoding=encoding, delimiter=delimiter, headerRow=2)
    stored = pyogrio.read_arrow(result["artifactPath"], layer="records")[1]
    assert stored["name"].to_pylist() == ["\u571f\u5730"]
    assert stored[result["dataset"]["sourceFidField"]].to_pylist() == ["3"]


@pytest.mark.parametrize("text", ["a,b\n1\n", "a,b\n1,2,3\n", 'a,b\n"unterminated,2\n'])
def test_malformed_csv_refused_without_snapshot(tmp_path, text):
    path = tmp_path / "invalid.csv"
    path.write_text(text, encoding="utf-8")
    with pytest.raises(DomainError):
        import_table(path, tmp_path / "work")
    assert not (tmp_path / "work" / "snapshot.gpkg").exists()


def test_csv_invalid_encoding_is_not_replaced(tmp_path):
    path = tmp_path / "invalid.csv"
    path.write_bytes(b"name\n\xff\n")
    with pytest.raises(DomainError) as error:
        import_table(path, tmp_path / "work")
    assert error.value.kind == "encoding_error"


def test_xlsx_enumeration_explicit_sheet_and_cell_provenance(sources, tmp_path):
    inspection = tables().inspect_table(options(sources["xlsx"], headerRow=2))
    assert inspection["sheets"] == ["coordinates", "other"]
    assert inspection["sheet"] is None
    assert inspection["columns"] == [] and inspection["rows"] == []
    with pytest.raises(InvalidParamsError):
        import_table(sources["xlsx"], tmp_path / "missing-sheet", headerRow=2)
    result = import_table(sources["xlsx"], tmp_path / "work", sheet="coordinates", headerRow=2)
    dataset = result["dataset"]
    assert dataset["cellMetadataLayer"] == "cell_metadata"
    stored = pyogrio.read_arrow(result["artifactPath"], layer="records")[1]
    assert stored["x"].to_pylist()[1] == "=114+1"
    assert stored["mixed"].to_pylist() == ["001", "true", "2026-09-09T12:30:00", "#DIV/0!", None]
    metadata = pyogrio.read_arrow(result["artifactPath"], layer="cell_metadata")[1].to_pylist()
    assert any(row["row_id"] == "1" and row["field_name"] == "x" and row["number_format"] == "000.00" for row in metadata)
    assert any(row["row_id"] == "2" and row["field_name"] == "x" and row["cell_type"] == "f" for row in metadata)
    assert dataset["source"]["metadata"]["formulaPolicy"] == "retain_text_no_execution"
    exported = vectors.export_vector({"dataset": dataset, "managedPath": result["artifactPath"]},
                                     tmp_path / "export", lambda *args: None, lambda: False)
    assert Path(exported["artifactPath"]).read_bytes() == Path(result["artifactPath"]).read_bytes()


def test_points_keep_invalid_rows_and_parent_identity(sources, tmp_path):
    table = import_table(sources["csv"], tmp_path / "table")
    before = Path(table["artifactPath"]).read_bytes()
    result = points(table, tmp_path / "points")
    dataset = result["dataset"]
    assert dataset["kind"] == "vector" and dataset["geometryType"] == "Point"
    assert "cellMetadataLayer" not in dataset
    assert dataset["featureCount"] == 4 and dataset["report"]["status"] == "restricted"
    assert dataset["report"]["counts"]["validCoordinates"] == 2
    assert dataset["report"]["counts"]["invalidCoordinates"] == 2
    metadata = dataset["source"]["metadata"]
    assert metadata["parentDatasetId"] == table["dataset"]["id"]
    assert metadata["parentVersion"] == table["dataset"]["version"]
    meta, stored = pyogrio.read_arrow(result["artifactPath"], layer="features")
    geometries = shapely.from_wkb(stored[meta["geometry_name"]].to_pylist())
    assert list(shapely.is_missing(geometries)) == [False, True, True, False]
    assert stored[dataset["sourceFidField"]].to_pylist() == ["2", "3", "4", "5"]
    assert stored[metadata["coordinateStatusField"]].to_pylist() == ["valid", "invalid", "invalid", "valid"]
    assert Path(table["artifactPath"]).read_bytes() == before
    vectors.verify_snapshot(Path(result["artifactPath"]), dataset)


def test_xlsx_formula_bool_and_date_coordinates_rejected(sources, tmp_path):
    table = import_table(sources["xlsx"], tmp_path / "table", sheet="coordinates", headerRow=2)
    result = points(table, tmp_path / "points")
    dataset = result["dataset"]
    assert dataset["report"]["counts"]["validCoordinates"] == 2
    assert dataset["report"]["counts"]["invalidCoordinates"] == 3
    meta, stored = pyogrio.read_arrow(result["artifactPath"], layer="features")
    reasons = stored[dataset["source"]["metadata"]["coordinateReasonField"]].to_pylist()
    assert "formula" in reasons[1] and "bool" in reasons[2] and "date" in reasons[3]


def test_projected_points_and_invalid_crs_policy(tmp_path):
    path = tmp_path / "projected.csv"
    path.write_text("x,y\n500000,3000000\n", encoding="utf-8")
    table = import_table(path, tmp_path / "table")
    result = points(table, tmp_path / "points", declaredCrs="EPSG:4547")
    assert result["dataset"]["crsAuthority"] == "EPSG:4547"
    assert result["dataset"]["bounds"] == [500000.0, 3000000.0, 500000.0, 3000000.0]
    assert result["dataset"]["boundsWgs84"] is not None
    for crs in (None, "EPSG:4978", "EPSG:4979"):
        with pytest.raises((DomainError, InvalidParamsError)):
            points(table, tmp_path / "invalid", declaredCrs=crs)


def test_all_invalid_coordinates_and_nonspatial_geometry_queries_refused(tmp_path):
    path = tmp_path / "invalid.csv"
    path.write_text("x,y\nNaN,27\nabc,0\n", encoding="utf-8")
    result = import_table(path, tmp_path / "table")
    with pytest.raises(DomainError) as error:
        points(result, tmp_path / "points")
    assert error.value.kind == "no_valid_coordinates"
    for query, params in [(vector_queries.viewport, {"bbox": [0, 0, 1, 1]}),
                          (vector_queries.feature, {"featureId": "1"})]:
        with pytest.raises(DomainError) as error:
            query(result["dataset"], Path(result["artifactPath"]), params)
        assert error.value.kind == "dataset_not_spatial"


def test_table_resource_limits_and_preview_budget(sources, tmp_path, monkeypatch):
    monkeypatch.setattr(tables(), "MAX_PREVIEW_ROWS", 2)
    preview = tables().inspect_table(options(sources["csv"]))
    assert len(preview["rows"]) == 2 and preview["truncated"]
    for limit_name, value in [("MAX_ROWS", 3), ("MAX_CELLS", 5), ("MAX_VALUE_BYTES", 10)]:
        with monkeypatch.context() as patch:
            patch.setattr(tables(), limit_name, value)
            with pytest.raises(DomainError) as error:
                import_table(sources["csv"], tmp_path / limit_name)
            assert error.value.kind == "source_limit"
    with monkeypatch.context() as patch:
        patch.setattr(tables(), "MAX_SOURCE_BYTES", 2)
        with pytest.raises(DomainError) as error:
            tables().inspect_table(options(sources["csv"]))
        assert error.value.kind == "source_limit"


def test_xlsx_zip_guard_and_table_cancellation(sources, tmp_path, monkeypatch):
    with monkeypatch.context() as patch:
        patch.setattr(tables(), "MAX_ZIP_EXPANDED_BYTES", 2)
        with pytest.raises(DomainError) as error:
            tables().inspect_table(options(sources["xlsx"]))
        assert error.value.kind == "source_limit"
    with pytest.raises(DomainError) as error:
        tables().import_table({**options(sources["csv"]), "datasetId": str(uuid.uuid4())}, tmp_path / "cancel",
                              progress=lambda *args: None, cancelled=lambda: True)
    assert error.value.kind == "task_cancelled"


def test_points_field_budget_reserves_internal_and_result_fields(tmp_path, monkeypatch):
    path = tmp_path / "many.csv"
    path.write_text("x,y,name\n114,27,one\n", encoding="utf-8")
    table = import_table(path, tmp_path / "table")
    monkeypatch.setattr(tables(), "MAX_FIELDS", 6)
    with pytest.raises(DomainError) as error:
        points(table, tmp_path / "points")
    assert error.value.kind == "source_limit"


def test_sparse_xlsx_late_columns_cannot_bypass_rectangular_cell_budget(tmp_path, monkeypatch):
    path = tmp_path / "sparse.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["x"])
    for _ in range(4):
        sheet.append(["one"])
    sheet.append(["one", "two", "three", "four"])
    workbook.save(path)
    workbook.close()
    monkeypatch.setattr(tables(), "MAX_CELLS", 12)
    with pytest.raises(DomainError) as error:
        import_table(path, tmp_path / "work", sheet="Sheet")
    assert error.value.kind == "source_limit"
    assert not (tmp_path / "work" / "snapshot.gpkg").exists()


@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_header_only_table_keeps_schema_and_cannot_generate_points(tmp_path, kind):
    path = tmp_path / f"empty.{kind}"
    if kind == "csv":
        path.write_text("x,y\n", encoding="utf-8")
        overrides = {}
    else:
        workbook = Workbook()
        workbook.active.append(["x", "y"])
        workbook.save(path)
        workbook.close()
        overrides = {"sheet": "Sheet"}
    result = import_table(path, tmp_path / "table", **overrides)
    assert result["dataset"]["featureCount"] == 0
    assert [field["name"] for field in result["dataset"]["fields"]] == ["x", "y"]
    with pytest.raises(DomainError) as error:
        points(result, tmp_path / "points")
    assert error.value.kind == "no_valid_coordinates"


def test_coordinate_numeric_grammar_is_ascii_decimal_or_scientific(tmp_path):
    path = tmp_path / "numbers.csv"
    path.write_text("x,y\n+1.14e2,2.7e1\n1_14,27\n\u0661\u0661\u0664,27\n1,234\n", encoding="utf-8")
    table = import_table(path, tmp_path / "table")
    result = points(table, tmp_path / "points")
    assert result["dataset"]["report"]["counts"]["validCoordinates"] == 1
    assert result["dataset"]["report"]["counts"]["invalidCoordinates"] == 3


@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_original_source_rows_survive_sort_filter_and_point_derivation(sources, tmp_path, kind):
    changes = {"sheet": "coordinates", "headerRow": 2} if kind == "xlsx" else {}
    imported = import_table(sources[kind], tmp_path / "table", **changes)
    derived = points(imported, tmp_path / "points")
    for result in (imported, derived):
        dataset = result["dataset"]
        page = vector_queries.attribute_page(dataset, Path(result["artifactPath"]), {
            "offset": 0, "limit": 500, "sortField": "code", "descending": True,
            "filter": {"field": "code", "operator": "contains", "value": "00"},
        })
        expected = [(str(index).zfill(3), str(index + changes.get("headerRow", 1)))
                    for index in range(dataset["featureCount"], 0, -1)]
        assert [(row["values"]["code"], row["sourceRow"]) for row in page["rows"]] == expected
    selected = vector_queries.feature(derived["dataset"], Path(derived["artifactPath"]), {"featureId": "2"})
    assert selected["row"]["sourceRow"] == str(2 + changes.get("headerRow", 1))


def test_xlsx_partial_preview_releases_the_source_handle(sources, tmp_path, monkeypatch):
    monkeypatch.setattr(tables(), "MAX_PREVIEW_ROWS", 1)
    inspection = tables().inspect_table(options(sources["xlsx"], sheet="coordinates", headerRow=2))
    assert inspection["truncated"]
    gc.collect()
    source = Path(sources["xlsx"])
    source.rename(tmp_path / "released.xlsx")
